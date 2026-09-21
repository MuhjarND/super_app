from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


DOCX = Path(r"C:\xampp\htdocs\super3\public\templates\surat-keterangan-perbaikan-presensi.docx")
KOP = Path(r"C:\xampp\htdocs\super3\public\kop_undangan.png")


def main() -> None:
    document = Document(DOCX)
    if not KOP.exists():
        raise FileNotFoundError(f"Kop surat tidak ditemukan: {KOP}")

    for section in document.sections:
        header = section.header
        paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        # The supplied letterhead is the official visual header for outgoing letters.
        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = 0
        paragraph.paragraph_format.space_after = 0
        paragraph.paragraph_format.line_spacing = 1
        paragraph.add_run().add_picture(str(KOP), width=Inches(6.30))
        section.header_distance = Inches(0.05)

    document.save(DOCX)
    print(f"Updated {DOCX}")


if __name__ == "__main__":
    main()
