from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from docx import Document


REFERENCE = Path(r"C:\Users\rubik\Documents\suket_perbaikan_presensi (2).docx")
OUTPUT = Path(r"C:\xampp\htdocs\super3\public\templates\surat-keterangan-perbaikan-presensi.docx")
EXPECTED_SHA256 = "71D50FBB3D520B80A746DC6C5A1D03242FCD22F32DC5D85E70AB8D216D7C30E9"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def replace_paragraph_text(paragraph, replacement: str) -> None:
    """Replace paragraph text while retaining the first run's direct formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(replacement)
        return
    runs[0].text = replacement
    for run in runs[1:]:
        run.text = ""


def iter_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
        yield from section.footer.paragraphs
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs


BODY_REPLACEMENTS = {
    2: "Nomor : {{nomor_surat}}",
    6: "Nama\t: {{nama_pembuat_keterangan}}",
    7: "NIP/NRP\t: {{nip_pembuat_keterangan}}",
    8: "Jabatan\t:\t{{jabatan_pembuat_keterangan}}",
    12: "Nama\t\t\t: {{nama_pegawai}}",
    13: "NIP/NRP\t\t\t: {{nip_pegawai}}",
    14: "Jabatan\t\t\t: {{jabatan_pegawai}}",
    15: "Unit Kerja\t\t\t: {{unit_kerja}}",
    16: "Satuan Kerja\t\t: {{satuan_kerja}}",
    17: "Tanggal Presensi\t: {{tanggal_presensi}}",
    18: "Hadir/pulang pukul\t: {{waktu_presensi}} {{zona_waktu}}",
    21: "Saya bertanggung jawab penuh atas kebenaran Informasi {{status_presensi}} nama tersebut di atas, sehubungan dengan hal tersebut mohon bantuannya untuk dilakukan perbaikan catatan jam kerja pada Sistem Informasi Manajemen Kepegawaian (SIKEP).",
    25: "{{tempat_surat}}, {{tanggal_surat}}",
    26: "{{jabatan_pembuat_keterangan}}",
    27: "{{tanda_tangan_pembuat_keterangan}}",
    28: "{{nama_pembuat_keterangan}}",
}

TABLE_REPLACEMENTS = {
    "Ditolak, karena": "Ditolak, karena {{alasan_penolakan}}",
    "<jabatan_pimpinan_satker>,": "{{jabatan_pimpinan_satker}},",
    "<tanda_tangan>": "{{tanda_tangan_pimpinan_satker}}",
    "<nama_pimpinan_satker>": "{{nama_pimpinan_satker}}",
    "NIP. <nip_pimpinan_satker>": "NIP. {{nip_pimpinan_satker}}",
}


def main() -> None:
    if sha256(REFERENCE) != EXPECTED_SHA256:
        raise RuntimeError("Reference DOCX hash changed; aborting to protect the source layout.")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    document = Document(OUTPUT)

    for index, replacement in BODY_REPLACEMENTS.items():
        paragraph = document.paragraphs[index]
        if not paragraph.text:
            raise RuntimeError(f"Expected source body paragraph {index} is empty.")
        replace_paragraph_text(paragraph, replacement)

    approval_cell = document.tables[0].rows[1].cells[0]
    found_table = set()
    for paragraph in approval_cell.paragraphs:
        original = paragraph.text
        key = next((candidate for candidate in TABLE_REPLACEMENTS if original == candidate or original.startswith(candidate)), None)
        replacement = TABLE_REPLACEMENTS.get(key) if key else None
        if replacement is not None:
            replace_paragraph_text(paragraph, replacement)
            found_table.add(key)
    missing = set(TABLE_REPLACEMENTS) - found_table
    if missing:
        raise RuntimeError("Expected approval paragraphs were not found: " + repr(sorted(missing)))

    document.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"SHA256 {sha256(OUTPUT)}")


if __name__ == "__main__":
    main()
